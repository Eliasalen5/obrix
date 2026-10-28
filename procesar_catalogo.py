import openpyxl
import csv
import re
import os
from datetime import datetime

inputFile = "C:/Users/HP 255R G10/Desktop/catalogoobrix.xlsx"
desktop = "C:/Users/HP 255R G10/Desktop"

wb = openpyxl.load_workbook(inputFile, data_only=True)
results = []
seen_key = set()
discarded = 0
duplicates_removed = 0

def norm(s):
    if s is None:
        return ''
    t = str(s).replace('\u00a0', ' ').replace('\t', ' ')
    t = ' '.join(t.split())
    return t.strip()

def norm_name(s):
    t = norm(s)
    if not t:
        return ''
    t2 = t
    t2 = re.sub(r'^Medida\s*:?\s*', '', t2, flags=re.I)
    t2 = re.sub(r'^C[oó]digo\s*:?\s*', '', t2, flags=re.I)
    for noise in ['Unidades por caja', 'Unidades de venta', 'Venta por', 'Caja cerrada', 'Unidad de venta']:
        t2 = re.sub(re.escape(noise), ' ', t2, flags=re.I)
    t2 = ' '.join(t2.split())
    return t2.strip()

def norm_meas(s):
    t = norm(s)
    if not t:
        return ''
    t2 = t.lower()
    t2 = t2.replace('\u00d7', 'x').replace('*', 'x')
    t2 = re.sub(r'^medida\s*:?\s*', '', t2)
    t2 = re.sub(r'\s*x\s*', ' x ', t2)
    t2 = ' '.join(t2.split())
    return t2.strip()

def is_noise(t):
    if not t:
        return True
    l = t.lower()
    if 'unidades por caja' in l: return True
    if 'unidades de venta' in l: return True
    if l.startswith('venta por'): return True
    if 'caja cerrada' in l: return True
    if l in ['medida', 'código', 'codigo', 'composición', 'composicion']: return True
    if re.match(r'^fila\s*\d+', l): return True
    if re.match(r'^hoja\s*', l): return True
    if len(t) < 3: return True
    return False

def is_code_like(s):
    t = norm(s)
    if re.match(r'^\d{4,6}$', t):
        return True
    return False

def meas_row(row):
    m = []
    for v in row:
        if v is None:
            continue
        t = norm(v)
        if not t:
            continue
        l = t.lower()
        if re.search(r'\d+\s*x\s*\d+', t):
            m.append(t)
            continue
        if re.search(r'\d+\s*(mm|cm|m|kg|lts|lt|gr|g|pulg)', t, re.I):
            m.append(t)
            continue
        if re.search(r'\d+/\d+', t) and any(x in l for x in ['mm', 'cm', 'm']):
            m.append(t)
            continue
        if re.search(r'\d+\s*-\s*\d+', t):
            m.append(t)
            continue
        if re.match(r'^\d+(\.\d+)?\s*(mm|cm|m|kg)$', t, re.I):
            m.append(t)
            continue
        if re.match(r'medida\s*:', l):
            m.append(re.sub(r'medida\s*:', '', t, flags=re.I).strip())
            continue
    sm = set()
    res = []
    for x in m:
        nm = norm_meas(x)
        if not nm or nm in sm:
            continue
        sm.add(nm)
        res.append(re.sub(r'medida\s*:', '', x, flags=re.I).strip())
    return ' | '.join(res)

def name_row(row, skip):
    best = ''
    for j, v in enumerate(row):
        if j == skip:
            continue
        if v is None:
            continue
        t = norm(v)
        if not t:
            continue
        if is_code_like(t):
            continue
        if is_noise(t):
            continue
        if re.match(r'^\d+$', t):
            continue
        if len(t) < 4:
            continue
        if len(t) > len(best):
            best = t
    if not best:
        return ''
    return norm_name(best)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        all_none = True
        for c in row:
            if c is not None:
                all_none = False
                break
        if all_none:
            continue
        ci = -1
        for j, v in enumerate(row):
            if is_code_like(v):
                ci = j
                break
        if ci == -1:
            discarded += 1
            continue
        name = name_row(row, ci)
        if not name or len(name) < 4 or is_noise(name):
            discarded += 1
            continue
        mr = meas_row(row)
        mn = norm_meas(mr)
        key = name.lower().replace(' ', '') + '||' + (mn.replace(' ', '') if mn else '')
        if key in seen_key:
            duplicates_removed += 1
            continue
        seen_key.add(key)
        results.append({'nombre': name, 'medidas': mr})

out = []
for i, r in enumerate(results):
    out.append({'sku': str(i + 1).zfill(5), 'nombre': r['nombre'], 'medidas': r['medidas']})

wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = 'Catalogo'
ws_out.append(['sku', 'nombre', 'medidas'])
for o in out:
    ws_out.append([o['sku'], o['nombre'], o['medidas']])
wb_out.save(os.path.join(desktop, 'catalogoobrix-sku-nuevos.xlsx'))

with open(os.path.join(desktop, 'catalogoobrix-sku-nuevos.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['sku', 'nombre', 'medidas'])
    for o in out:
        w.writerow([o['sku'], o['nombre'], o['medidas']])

rpt = []
rpt.append('CATALOGO OBRIX - SKU NUEVOS')
rpt.append('===========================')
rpt.append('Fecha: ' + datetime.now().isoformat())
rpt.append('')
rpt.append('Reglas: SKU desde 00001, duplicados por clave compuesta (nombre+medidas normalizados), primera aparicion')
rpt.append('')
rpt.append('Productos finales: ' + str(len(out)))
rpt.append('Descartados (sin nombre claro/ruido): ' + str(discarded))
rpt.append('Duplicados eliminados: ' + str(duplicates_removed))
rpt.append('')
rpt.append('Primeros 30:')
for i, o in enumerate(out[:30]):
    rpt.append(str(i + 1) + '. ' + o['sku'] + ' | ' + o['nombre'] + (' | ' + o['medidas'] if o['medidas'] else ''))
with open(os.path.join(desktop, 'catalogoobrix-sku-nuevos-reporte.txt'), 'w', encoding='utf-8') as f:
    f.write('\r\n'.join(rpt) + '\r\n')

print('OK')
print('Productos:', len(out))
print('Descartados:', discarded)
print('Duplicados eliminados:', duplicates_removed)
print('Archivos en escritorio:')
print('- catalogoobrix-sku-nuevos.xlsx')
print('- catalogoobrix-sku-nuevos.csv')
print('- catalogoobrix-sku-nuevos-reporte.txt')
